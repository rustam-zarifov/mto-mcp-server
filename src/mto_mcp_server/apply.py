"""apply_mto_update: re-runs the diff, then (optionally) writes it.

Deliberately does NOT accept a pre-computed diff object from the caller -
it recomputes from the same two file paths every time. That keeps the
tool safe against a stale/edited diff being replayed, and keeps the two
MCP tools independent: diff_mto_scsa is always a pure preview, apply
always re-derives ground truth right before writing.

Multi-row redistribution (`_redistribute_group`) is ported from
`redistribute_group()` in
https://github.com/rustam-zarifov/mto-scsa-reconciliation's
`5_apply_fixes.py`: when one match key is backed by several MTO rows
(e.g. the same elbow spec appearing on two different lines), scale each
row's quantity proportionally to hit the new total rather than dumping
everything into one row and zeroing the rest.
"""
from __future__ import annotations

import openpyxl

from .diff import diff_mto_scsa
from .io_excel import MTO_COLUMNS

_DEFAULT_UNIT_BY_CATEGORY = {"pipe": "м"}
_DEFAULT_UNIT = "шт"


def _redistribute_group(rows_with_qty: list[tuple[int, float]], target_total: float) -> dict:
    """rows_with_qty: [(row_index, current_qty), ...] for one match key.

    Scales each row's current qty proportionally to hit target_total,
    rounds to integers, then pushes any rounding remainder onto the row
    with the largest current quantity. If current total is 0, splits the
    target evenly instead. Returns {row_index: new_qty}.
    """
    current_total = sum(q for _, q in rows_with_qty)
    n = len(rows_with_qty)

    if current_total == 0:
        base = int(target_total) // n
        remainder = int(target_total) - base * n
        return {
            row_idx: base + (1 if i < remainder else 0)
            for i, (row_idx, _) in enumerate(rows_with_qty)
        }

    allocated = [round(q / current_total * target_total) for _, q in rows_with_qty]
    diff = target_total - sum(allocated)
    idx_max = max(range(n), key=lambda i: rows_with_qty[i][1])
    allocated[idx_max] += diff
    return {
        rows_with_qty[i][0]: max(allocated[i], 0) for i in range(n)
    }


def apply_mto_update(
    mto_path: str,
    scsa_path: str,
    category: str | None = None,
    include_removals: bool = False,
    dry_run: bool = False,
    output_path: str | None = None,
) -> dict:
    result = diff_mto_scsa(mto_path, scsa_path, category=category)

    write_target = output_path or mto_path
    writes = {"qty_updated": 0, "zeroed": 0, "inserted": 0}

    if dry_run:
        result["applied"] = False
        result["dry_run"] = True
        result["would_write"] = {
            "qty_updated": len(result["qty_mismatch"]),
            "zeroed": len(result["removed"]) if include_removals else 0,
            "inserted": len(result["new"]),
        }
        result["include_removals"] = include_removals
        return result

    wb = openpyxl.load_workbook(mto_path)
    ws = wb.active
    qty_col = MTO_COLUMNS["qty"] + 1  # openpyxl is 1-indexed

    # Scenario 1: qty mismatch -> single row gets a direct overwrite;
    # multiple rows sharing the key get redistribute_group()'s proportional
    # split instead of dumping the total into one row and zeroing the rest.
    for item in result["qty_mismatch"]:
        rows = item["mto_row_indices"]
        if len(rows) == 1:
            ws.cell(row=rows[0] + 1, column=qty_col, value=item["scsa_qty"])
        else:
            rows_with_qty = [
                (r, ws.cell(row=r + 1, column=qty_col).value or 0) for r in rows
            ]
            fixes = _redistribute_group(rows_with_qty, item["scsa_qty"])
            for r, new_qty in fixes.items():
                ws.cell(row=r + 1, column=qty_col, value=new_qty)
        writes["qty_updated"] += 1

    # Scenario 2: removed -> zero out, but only if explicitly requested.
    if include_removals:
        for item in result["removed"]:
            for r in item["mto_row_indices"]:
                ws.cell(row=r + 1, column=qty_col, value=0)
            writes["zeroed"] += 1

    # Scenario 3: new -> append a row at the bottom.
    if result["new"]:
        existing_poz = [
            row[0].value
            for row in ws.iter_rows(min_col=1, max_col=1)
            if isinstance(row[0].value, (int, float))
        ]
        next_poz = int(max(existing_poz)) + 1 if existing_poz else 1
        for item in result["new"]:
            unit = _DEFAULT_UNIT_BY_CATEGORY.get(item["category"], _DEFAULT_UNIT)
            ws.append([next_poz, item["name"], item["scsa_qty"], unit, item["line"]])
            next_poz += 1
            writes["inserted"] += 1

    wb.save(write_target)

    result["applied"] = True
    result["dry_run"] = False
    result["include_removals"] = include_removals
    result["output_path"] = write_target
    result["writes"] = writes
    return result
