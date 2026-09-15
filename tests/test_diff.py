"""Tests run against real (small) MTO/SCSA export pairs in tests/fixtures.

Expected numbers were derived by hand-tracing the fixture files, not by
running the code and asserting whatever it happened to output - see the
comment above each fixture-derived value.
"""
import shutil
from pathlib import Path

import openpyxl
import pytest

from mto_mcp_server.apply import apply_mto_update
from mto_mcp_server.diff import diff_mto_scsa

FIXTURES = Path(__file__).parent / "fixtures"
MTO = str(FIXTURES / "mto_test.xlsx")
SCSA = str(FIXTURES / "scsa_test.xlsx")


@pytest.fixture()
def result():
    return diff_mto_scsa(MTO, SCSA)


def test_exact_match_is_not_flagged(result):
    # "Отвод 90 FRP DN20" / N1A: MTO qty=2, SCSA has 2 physical rows.
    # Must count as matched, must NOT appear in qty_mismatch.
    assert result["summary"]["matched"] == 1
    names_in_mismatch = {r["name"] for r in result["qty_mismatch"]}
    assert "Отвод 90 FRP DN20" not in names_in_mismatch


def test_qty_mismatch_both_directions(result):
    by_name = {r["name"]: r for r in result["qty_mismatch"]}
    # SCSA qty (4) > MTO qty (3) - increase
    assert by_name["Отвод 90 DN80 Сталь20"]["mto_qty"] == 3
    assert by_name["Отвод 90 DN80 Сталь20"]["scsa_qty"] == 4
    # SCSA qty (2) < MTO qty (1+3=4, summed across 2 MTO rows) - decrease
    assert by_name["Отвод 90 FRP DN1400"]["mto_qty"] == 4
    assert by_name["Отвод 90 FRP DN1400"]["scsa_qty"] == 2


def test_removed_items(result):
    names = {r["name"] for r in result["removed"]}
    assert names == {"Отвод 45 FRP DN150", "Отвод 90 DN50 Сталь20"}


def test_new_items(result):
    names = {r["name"] for r in result["new"]}
    assert names == {"Отвод 90 FRP DN1200", "Отвод 90 FRP DN150"}


def test_dry_run_writes_nothing(tmp_path):
    scratch = tmp_path / "mto.xlsx"
    shutil.copy(MTO, scratch)
    before = scratch.read_bytes()
    apply_mto_update(str(scratch), SCSA, dry_run=True, include_removals=True)
    assert scratch.read_bytes() == before


def test_apply_without_removals_leaves_removed_items_untouched(tmp_path):
    scratch = tmp_path / "mto.xlsx"
    shutil.copy(MTO, scratch)
    apply_mto_update(str(scratch), SCSA, dry_run=False, include_removals=False)
    ws = openpyxl.load_workbook(scratch).active
    rows = {r[1]: r[2] for r in ws.iter_rows(values_only=True)}
    # was qty=1 in MTO, absent from SCSA - must survive untouched
    assert rows["Отвод 45 FRP DN150"] == 1
    assert rows["Отвод 90 DN50 Сталь20"] == 1


def test_apply_with_removals_zeroes_them(tmp_path):
    scratch = tmp_path / "mto.xlsx"
    shutil.copy(MTO, scratch)
    apply_mto_update(str(scratch), SCSA, dry_run=False, include_removals=True)
    ws = openpyxl.load_workbook(scratch).active
    rows = {r[1]: r[2] for r in ws.iter_rows(values_only=True)}
    assert rows["Отвод 45 FRP DN150"] == 0
    assert rows["Отвод 90 DN50 Сталь20"] == 0


def test_new_items_get_scsa_qty_and_next_poz(tmp_path):
    scratch = tmp_path / "mto.xlsx"
    shutil.copy(MTO, scratch)
    apply_mto_update(str(scratch), SCSA, dry_run=False)
    ws = openpyxl.load_workbook(scratch).active
    rows = list(ws.iter_rows(values_only=True))
    last_two_names = {r[1] for r in rows[-2:]}
    assert last_two_names == {"Отвод 90 FRP DN1200", "Отвод 90 FRP DN150"}
    assert all(r[2] == 1 for r in rows[-2:])


def test_qty_mismatch_multi_row_redistributes_proportionally(tmp_path):
    # "Отвод 90 FRP DN1400" is backed by 2 MTO rows (qty 1 on line N1A,
    # qty 3 on line L1A; total 4) but SCSA only shows 2 physical
    # instances. redistribute_group() scales proportionally instead of
    # collapsing everything onto one row: 1/4*2=0.5->0 (banker's
    # rounding), 3/4*2=1.5->2, remainder 0 - so N1A -> 0, L1A -> 2.
    scratch = tmp_path / "mto.xlsx"
    shutil.copy(MTO, scratch)
    apply_mto_update(str(scratch), SCSA, dry_run=False)
    ws = openpyxl.load_workbook(scratch).active
    by_line = {
        r[4]: r[2]
        for r in ws.iter_rows(values_only=True)
        if r[1] == "Отвод 90 FRP DN1400"
    }
    assert by_line == {"N1A": 0, "L1A": 2}
    assert sum(by_line.values()) == 2  # matches SCSA's total
